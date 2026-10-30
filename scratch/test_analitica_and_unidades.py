# scratch/test_analitica_and_unidades.py
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import json
import openpyxl
from io import BytesIO

def test_unidades_metricas():
    print("--- Probando rutas_unidades_metricas ---")
    from rutas_unidades_metricas import obtener_metricas_flota
    flota = obtener_metricas_flota()
    totales = flota["totales"]
    porcentajes = flota["porcentajes"]
    estados = flota["estados_por_unidad"]
    print(f"Totales: {totales}")
    print(f"Porcentajes: {porcentajes}")
    
    # Check if any unit is in 'Garantía/Retrabajo'
    unidades_garantia = [u for u, est in estados.items() if est == "Garantía/Retrabajo"]
    print(f"Unidades en 'Garantía/Retrabajo': {unidades_garantia}")
    assert "garantia" in totales, "Debe existir key 'garantia' en totales"
    assert "garantia" in porcentajes, "Debe existir key 'garantia' en porcentajes"
    print(">> Metricas de flota OK")

def test_analitica_auditada():
    print("\n--- Probando rutas_analitica.obtener_datos_auditados ---")
    from rutas_analitica import obtener_datos_auditados
    datos = obtener_datos_auditados("Todos")
    
    events = datos["events"]
    print(f"Total eventos registrados: {len(events)}")
    assert "ultima_garantia" in events, "Debe existir 'ultima_garantia' en events"
    ug = events["ultima_garantia"]
    print(f"Última unidad a garantía: {ug}")
    if ug:
        print(f"  Ticket: {ug.get('ticket')}")
        print(f"  Unidad: {ug.get('unidad')}")
        print(f"  Taller: {ug.get('taller')}")
        print(f"  Supervisor: {ug.get('supervisor')}")
        print(f"  Fecha: {ug.get('fecha')}")
    
    charts = datos["charts"]
    assert "pie_flota" in charts, "Debe existir pie_flota en charts"
    pie = charts["pie_flota"]
    print(f"Pie Flota Labels: {pie['labels']}")
    print(f"Pie Flota Data: {pie['data']}")
    assert "Garantía/Retrabajo" in pie["labels"], "'Garantía/Retrabajo' debe estar en los labels de pie_flota"
    assert len(pie["labels"]) == 5, f"Deben ser 5 categorías en pie_flota, hay {len(pie['labels'])}"
    assert len(pie["data"]) == 5, f"Deben ser 5 valores en pie_flota, hay {len(pie['data'])}"
    print(">> Datos auditados y pie_flota OK")

def test_flask_endpoints():
    print("\n--- Probando Flask Client ---")
    from app import app
    client = app.test_client()
    
    # 1. /api/unidades
    with client.session_transaction() as sess:
        sess["usuario"] = {"nombre": "Admin Test", "rol": "administracion"}
    
    res = client.get("/api/unidades")
    assert res.status_code == 200, f"Status code {res.status_code}"
    unidades = res.get_json()
    print(f"Total unidades devueltas por /api/unidades: {len(unidades)}")
    unidades_gar = [k for k, v in unidades.items() if v.get("estado_operativo") == "Garantía/Retrabajo"]
    print(f"Unidades con estado_operativo == 'Garantía/Retrabajo': {unidades_gar}")
    
    # 2. /api/analitica/ultimos
    res_ult = client.get("/api/analitica/ultimos?period=Todos")
    assert res_ult.status_code == 200, f"Status code {res_ult.status_code}"
    ult_json = res_ult.get_json()
    assert "ultima_garantia" in ult_json, "Debe incluir 'ultima_garantia'"
    assert "flota_metricas" in ult_json, "Debe incluir 'flota_metricas'"
    print(f"Ultima garantia API: {ult_json.get('ultima_garantia')}")
    
    # 3. /api/analitica/charts
    res_charts = client.get("/api/analitica/charts?period=Todos")
    assert res_charts.status_code == 200
    charts_json = res_charts.get_json()
    assert "pie_flota" in charts_json
    assert "Garantía/Retrabajo" in charts_json["pie_flota"]["labels"]
    print(">> Endpoints /api/unidades, /api/analitica/ultimos y /api/analitica/charts OK")

    # 4. /api/analitica/exportar_excel
    res_excel = client.get("/api/analitica/exportar_excel?period=Todos")
    assert res_excel.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(res_excel.data))
    assert "Eventos Clave" in wb.sheetnames
    assert "Distribución y Gráficas" in wb.sheetnames
    
    ws1 = wb["Eventos Clave"]
    found_garantia_event = False
    for r in range(4, 20):
        val = ws1.cell(row=r, column=1).value
        if val and "Garantía" in str(val):
            found_garantia_event = True
            print(f"Fila encontrada en Excel Hoja 1: {val} -> {ws1.cell(row=r, column=2).value} (Eco: {ws1.cell(row=r, column=4).value})")
            break
    assert found_garantia_event, "No se encontró el evento de Garantía en la Hoja 1 de Excel"

    ws3 = wb["Distribución y Gráficas"]
    found_pie_flota_garantia = False
    for r in range(1, 50):
        val = ws3.cell(row=r, column=1).value
        if val and "Garantía/Retrabajo" in str(val):
            found_pie_flota_garantia = True
            print(f"Fila encontrada en Excel Hoja 3: {val} -> Cantidad: {ws3.cell(row=r, column=2).value}")
            break
    assert found_pie_flota_garantia, "No se encontró Garantía/Retrabajo en Hoja 3 de Excel"
    print(">> Exportación a Excel verificada al 100%")

if __name__ == "__main__":
    test_unidades_metricas()
    test_analitica_auditada()
    test_flask_endpoints()
    print("\n==========================================")
    print("TODAS LAS PRUEBAS PASARON EXITOSAMENTE (100%)")
    print("==========================================")
