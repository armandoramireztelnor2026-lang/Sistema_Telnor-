from flask import Blueprint, jsonify, request, send_file
import json
import os
import datetime
from collections import Counter
import re
import io

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from rutas_seccion_reportes import determinar_estado_ticket

analitica_bp = Blueprint("analitica_bp", __name__)

def leer_json(archivo):
    if not os.path.exists(archivo): return {}
    try:
        with open(archivo, 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        try:
            with open(archivo, 'r', encoding='latin-1') as f:
                return json.load(f)
        except Exception:
            return {}

def obtener_ticket_id_factura(f):
    if not f: return ''
    for field in ['ticket', 'ticket_id', 'id_reporte', 'numero_reporte', 'id_ticket', 'reporte_id']:
        val = f.get(field)
        if val and str(val).strip() and str(val).strip().upper() != 'N/A':
            return str(val).strip()
    match = re.search(r"\[TICKET:\s*([A-Za-z0-9_-]+)\]", str(f.get('retro', '')))
    if match:
        return match.group(1).strip()
    return ''

def limpiar_unidad(u_str):
    if not u_str: return ''
    s = str(u_str).strip()
    if '-' in s:
        s = s.split('-')[-1].strip()
    return s

def obtener_unidad(f=None, r=None):
    if r and r.get('unidad'):
        return limpiar_unidad(r.get('unidad'))
    if f:
        if f.get('unidad_eco'):
            return limpiar_unidad(f.get('unidad_eco'))
        if f.get('unidad'):
            return limpiar_unidad(f.get('unidad'))
    return '1000'

def obtener_chofer(f=None, r=None):
    if r:
        e = r.get('empleado') or r.get('chofer') or r.get('nombre_empleado')
        if e and str(e).strip().upper() not in ('', 'N/A', 'NONE', 'NULL'):
            return str(e).strip()
    if f:
        ch = f.get('chofer_nombre') or f.get('chofer') or f.get('responsable') or f.get('empleado')
        if ch and str(ch).strip().upper() not in ('', 'N/A', 'NONE', 'NULL'):
            return str(ch).strip()
        match = re.search(r"CHOFER:\s*([^\r\n]+)", str(f.get('retro', '')))
        if match:
            return match.group(1).strip()
    return 'Personal Operativo'

def obtener_taller(f=None, r=None):
    if f:
        t = f.get('proveedor') or f.get('taller')
        if t and str(t).strip().upper() not in ('', 'N/A', 'NONE', 'NULL'):
            return str(t).strip()
    if r:
        t = r.get('asignado_a')
        if t and str(t).strip().upper() not in ('', 'N/A', 'NONE', 'NULL'):
            return str(t).strip()
    return 'AUTOPARTES Y MOTORES B.C.'

def obtener_todas_facturas():
    facturas_raw = leer_json('facturas.json').get('facturas', [])
    archivados_raw = leer_json('tickets_archivados.json').get('facturas', [])
    
    facturas_dict = {}
    for f in facturas_raw:
        if f.get('id'):
            facturas_dict[str(f.get('id'))] = f
    for f in archivados_raw:
        if f.get('id') and str(f.get('id')) not in facturas_dict:
            facturas_dict[str(f.get('id'))] = f
            
    facturas_all = list(facturas_dict.values())
    
    def fecha_sort_key(item):
        return str(item.get('timestamp') or item.get('timestamp_aprobacion') or item.get('fecha_cierre') or item.get('fecha') or '')
        
    facturas_all.sort(key=fecha_sort_key)
    return facturas_all

def construir_mapa_responsables(usuarios_data):
    cope_map = {}
    default_admin = 'Manuel Alejandro Marquez'
    default_super = 'Samuel Giron'
    default_jefe = 'Jose Alejandro Peñuñuri'

    for u in usuarios_data:
        if u.get('rol') == 'administracion':
            dp = u.get('datos_perfil', {})
            cope = dp.get('cope')
            subrol = dp.get('subrol', '').capitalize()
            nombres = dp.get('nombres', '')
            apellidos = dp.get('apellido_paterno', '')
            name = re.sub(r'\s+', ' ', f"{nombres} {apellidos}").strip()
            
            if subrol == 'Administrador' and name: default_admin = name
            if subrol == 'Supervisor' and name: default_super = name
            if subrol == 'Jefatura' and name: default_jefe = name

            copes = [cope] if cope else []
            copes.extend(dp.get('copes_asignados', []))
            for c in copes:
                if not c: continue
                c_key = c.strip().upper()
                if c_key not in cope_map:
                    cope_map[c_key] = {}
                cope_map[c_key][subrol] = name
                
    return cope_map, default_admin, default_super, default_jefe

def resolver_responsables(cope, cope_map, default_admin, default_super, default_jefe):
    if not cope:
        cope_k = 'PE CALLE 10'
    else:
        cope_k = cope.strip().upper()
    info = cope_map.get(cope_k, {})
    return {
        'Administrador': info.get('Administrador') or default_admin,
        'Supervisor': info.get('Supervisor') or default_super,
        'Jefatura': info.get('Jefatura') or default_jefe
    }

def filtrar_por_periodo(reportes, facturas, period_str):
    if not period_str or period_str == 'Todos' or period_str == '':
        return reportes, facturas
        
    parts = period_str.split('-')
    target_year = parts[0]
    target_semester = parts[1] if len(parts) > 1 else None
        
    rep_filtrados = []
    ticket_periods = {}
    
    for r in reportes:
        y = str(r.get('fecha', ''))[:4]
        m = str(r.get('fecha', ''))[5:7]
        
        if not y and r.get('timestamp'):
            y = str(r.get('timestamp', ''))[:4]
            m = str(r.get('timestamp', ''))[5:7]
            
        ticket_id = str(r.get('id', ''))
        
        try:
            month_int = int(m)
            semester = 'S1' if month_int <= 6 else 'S2'
        except ValueError:
            semester = 'S1'
            
        ticket_periods[ticket_id] = (y, semester)
        
        if target_semester:
            if y == target_year and semester == target_semester:
                rep_filtrados.append(r)
        else:
            if y == target_year:
                rep_filtrados.append(r)
            
    fact_filtradas = []
    for f in facturas:
        ticket_id = obtener_ticket_id_factura(f)
        y, semester = ticket_periods.get(ticket_id, (None, None))
        
        if not y:
            f_date = str(f.get('fecha') or f.get('timestamp') or f.get('timestamp_aprobacion') or f.get('fecha_cierre') or '')
            if len(f_date) >= 7:
                y = f_date[:4]
                m = f_date[5:7]
                try:
                    semester = 'S1' if int(m) <= 6 else 'S2'
                except ValueError:
                    semester = 'S1'
                    
        if target_semester:
            if y == target_year and semester == target_semester:
                fact_filtradas.append(f)
        else:
            if y == target_year:
                fact_filtradas.append(f)
            
    return rep_filtrados, fact_filtradas

def obtener_datos_auditados(period):
    usuarios_data = leer_json('usuarios.json').get('usuarios', [])
    reportes_raw = leer_json('reportes.json').get('reportes', [])
    facturas_all = obtener_todas_facturas()

    def fecha_sort_key(item):
        return str(item.get('timestamp') or item.get('timestamp_aprobacion') or item.get('fecha_cierre') or item.get('fecha') or '')
    reportes_raw.sort(key=fecha_sort_key)

    reportes, facturas = filtrar_por_periodo(reportes_raw, facturas_all, period)
    reportes_map = {str(r.get('id')): r for r in reportes_raw if r.get('id')}

    cope_map, def_admin, def_super, def_jefe = construir_mapa_responsables(usuarios_data)

    def extract_factura_info(f):
        tid = obtener_ticket_id_factura(f)
        r = reportes_map.get(tid)
        
        chofer = obtener_chofer(f, r)
        taller = obtener_taller(f, r)
        unidad = obtener_unidad(f, r)
        cope = (r.get('cope') if r else None) or f.get('cope_origen') or f.get('cope') or 'PE CALLE 10'
        resp = resolver_responsables(cope, cope_map, def_admin, def_super, def_jefe)
        
        fecha = f.get('fecha_cierre') or f.get('timestamp_aprobacion') or f.get('timestamp') or f.get('fecha')
        if not fecha and r:
            fecha = r.get('timestamp') or r.get('fecha')
        if not fecha:
            fecha = datetime.datetime.now().isoformat()
            
        return {
            "ticket": tid or f.get('id', 'N/A'),
            "chofer": chofer,
            "unidad": unidad,
            "taller": taller,
            "administrador": resp['Administrador'],
            "supervisor": resp['Supervisor'],
            "jefatura": resp['Jefatura'],
            "fecha": fecha
        }

    # 1. Ultimo llegado
    ultimo_llegado = None
    if reportes:
        r = reportes[-1]
        resp = resolver_responsables(r.get('cope'), cope_map, def_admin, def_super, def_jefe)
        ultimo_llegado = {
            "ticket": r.get('id', 'N/A'),
            "chofer": obtener_chofer(None, r),
            "unidad": obtener_unidad(None, r),
            "taller": obtener_taller(None, r),
            "administrador": resp['Administrador'],
            "supervisor": resp['Supervisor'],
            "jefatura": resp['Jefatura'],
            "fecha": r.get('timestamp') or r.get('fecha', '')
        }

    # 2. Ultimo finalizado
    ultimo_finalizado = None
    for f in reversed(facturas):
        if f.get('estado') in ['Archivado', 'Archivado_Final'] or (f.get('liberado_admin') is True and (f.get('factura_pdf') or f.get('factura_folio'))) or f.get('fecha_cierre'):
            ultimo_finalizado = extract_factura_info(f)
            break

    # 3. Ultimo asignado
    ultimo_asignado = None
    if facturas:
        ultimo_asignado = extract_factura_info(facturas[-1])
    elif reportes:
        ultimo_asignado = extract_factura_info({'id': reportes[-1].get('id'), 'retro': f"[TICKET:{reportes[-1].get('id')}]"})

    # 4. Ultimo caro
    ultimo_caro = None
    for f in reversed(facturas):
        p = float(f.get('precio_estimado', f.get('precio', 0)) or 0)
        if p >= 10001:
            ultimo_caro = extract_factura_info(f)
            break
    if not ultimo_caro and facturas:
        f_max = max(facturas, key=lambda x: float(x.get('precio_estimado', x.get('precio', 0)) or 0))
        ultimo_caro = extract_factura_info(f_max)

    # 5. Ultima reparada
    ultima_reparada = None
    for f in reversed(facturas):
        if f.get('estado') in ['Reparado', 'Archivado', 'Archivado_Final'] or f.get('entregado') in ['Sí', 'S'] or f.get('codigo_liberacion'):
            ultima_reparada = extract_factura_info(f)
            break

    # 6. Ultima entregada
    ultima_entregada = None
    for f in reversed(facturas):
        if f.get('entregado') in ['Sí', 'S'] or f.get('codigo_liberacion') or f.get('estado') in ['Archivado', 'Archivado_Final']:
            ultima_entregada = extract_factura_info(f)
            break

    # 7. Ultima rechazada
    ultima_rechazada = None
    for f in reversed(facturas):
        st_f = str(f.get('estado', '')).lower()
        retro_f = str(f.get('retro', '')).lower()
        has_rej_cot = any('rechazad' in str(c.get('estado', '')).lower() or 'cancelad' in str(c.get('estado', '')).lower() for c in f.get('cotizaciones', []))
        if 'rechazad' in st_f or 'cancelad' in st_f or 'rechazad' in retro_f or 'cancelad' in retro_f or has_rej_cot:
            ultima_rechazada = extract_factura_info(f)
            break
    if not ultima_rechazada and reportes:
        for r in reversed(reportes):
            st_r = str(r.get('estado', '')).lower()
            if 'rechazad' in st_r or 'cancelad' in st_r:
                resp = resolver_responsables(r.get('cope'), cope_map, def_admin, def_super, def_jefe)
                ultima_rechazada = {
                    "ticket": r.get('id'),
                    "chofer": obtener_chofer(None, r),
                    "unidad": obtener_unidad(None, r),
                    "taller": obtener_taller(None, r),
                    "administrador": resp['Administrador'],
                    "supervisor": resp['Supervisor'],
                    "jefatura": resp['Jefatura'],
                    "fecha": r.get('timestamp') or r.get('fecha')
                }
                break

    # 8. Ultima fiscal
    ultima_fiscal = None
    for f in reversed(facturas):
        has_cot_fol = any(bool(c.get('factura_folio')) for c in f.get('cotizaciones', []))
        if bool(f.get('factura_folio')) or bool(f.get('factura_pdf')) or bool(f.get('pdf_fiscal')) or bool(f.get('numero_doc50')) or f.get('validacion_fiscal') in ['Aprobada', 'Pendiente', 'Revisión'] or has_cot_fol:
            ultima_fiscal = extract_factura_info(f)
            break

    # 9. Ultimo ticket reactivado de cotizaciones caras
    ultimo_reactivado_caras = None
    candidatos_reactivados = []
    for f in facturas:
        tid = obtener_ticket_id_factura(f)
        rep = reportes_map.get(tid)
        if f.get('reactivado_desde_caras') is True or (rep and rep.get('reactivado_desde_caras') is True):
            candidatos_reactivados.append(extract_factura_info(f))

    for r in reportes:
        if r.get('reactivado_desde_caras') is True:
            r_id = str(r.get('id', ''))
            if not any(str(c.get('ticket')) == r_id for c in candidatos_reactivados):
                resp = resolver_responsables(r.get('cope'), cope_map, def_admin, def_super, def_jefe)
                candidatos_reactivados.append({
                    "ticket": r.get('id', 'N/A'),
                    "chofer": obtener_chofer(None, r),
                    "unidad": obtener_unidad(None, r),
                    "taller": obtener_taller(None, r),
                    "administrador": resp['Administrador'],
                    "supervisor": resp['Supervisor'],
                    "jefatura": resp['Jefatura'],
                    "fecha": r.get('timestamp') or r.get('fecha', '')
                })

    if candidatos_reactivados:
        candidatos_reactivados.sort(key=lambda x: str(x.get('fecha') or ''))
        ultimo_reactivado_caras = candidatos_reactivados[-1]

    # 10. Ultima garantia reasignada (supervisor reasigna un trabajo a un taller)
    retrabajos_raw = leer_json('retrabajos.json').get('retrabajos', [])
    candidatos_garantia = []
    for ret in retrabajos_raw:
        st_ret = str(ret.get('estado', '')).strip().lower()
        f_reasig = ret.get('fecha_reasignacion') or ''
        # Se considera si fue reasignado o tiene fecha de reasignacion registrada
        if st_ret in ['reasignado', 'en proceso', 'finalizado'] or f_reasig:
            fecha_ref = f_reasig or ret.get('timestamp') or ret.get('fecha_solicitud') or ''
            if period != 'Todos' and fecha_ref:
                parts = period.split('-')
                target_year = parts[0]
                target_sem = parts[1] if len(parts) > 1 else None
                y_ref = str(fecha_ref)[:4]
                m_ref = str(fecha_ref)[5:7] if len(str(fecha_ref)) >= 7 else '01'
                try:
                    sem_ref = 'S1' if int(m_ref) <= 6 else 'S2'
                except Exception:
                    sem_ref = 'S1'
                if target_sem:
                    if y_ref != target_year or sem_ref != target_sem:
                        continue
                else:
                    if y_ref != target_year:
                        continue

            cope_val = ret.get('cope') or (ret.get('seccion_1', {}).get('cope')) or 'PE CALLE 10'
            resp = resolver_responsables(cope_val, cope_map, def_admin, def_super, def_jefe)
            sup_nombre = ret.get('supervisor_reasigno') or ret.get('supervisor_asignado') or resp['Supervisor']
            chofer_nombre = ret.get('empleado') or (ret.get('seccion_1', {}).get('empleado')) or ret.get('chofer') or 'Personal Operativo'
            taller_nombre = ret.get('proveedor') or (ret.get('seccion_2', {}).get('proveedor')) or 'Taller'

            candidatos_garantia.append({
                "ticket": ret.get('ticket_original') or ret.get('ticket') or ret.get('id', 'N/A'),
                "chofer": chofer_nombre,
                "unidad": str(ret.get('unidad', '')).replace('8090-', ''),
                "taller": taller_nombre,
                "administrador": resp['Administrador'],
                "supervisor": sup_nombre,
                "jefatura": resp['Jefatura'],
                "fecha": fecha_ref or datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "motivo": ret.get('motivo_chofer', '')
            })

    ultima_garantia = None
    if candidatos_garantia:
        candidatos_garantia.sort(key=lambda x: str(x.get('fecha') or ''))
        ultima_garantia = candidatos_garantia[-1]

    # TOPS
    now = datetime.datetime.now()
    unidades_tiempo = []
    for f in facturas:
        st_f = f.get('estado')
        entregado = f.get('entregado') in ['Sí', 'S']
        if st_f not in ['Archivado', 'Archivado_Final', 'Cancelado_Cotizacion_Cara', 'Rechazado', 'Eliminado'] and not entregado:
            dt = now
            ts = f.get('timestamp_aprobacion') or f.get('timestamp') or f.get('fecha')
            if ts:
                try:
                    dt = datetime.datetime.fromisoformat(str(ts)[:19])
                except Exception:
                    try:
                        dt = datetime.datetime.strptime(str(ts)[:10], "%Y-%m-%d")
                    except Exception:
                        dt = now
            days = (now - dt).days
            if days < 0: days = 0
            
            tid = obtener_ticket_id_factura(f)
            r = reportes_map.get(tid)
            u = obtener_unidad(f, r)
            t = obtener_taller(f, r)
            unidades_tiempo.append({"unidad": u, "dias": days, "taller": t})

    unidades_tiempo.sort(key=lambda x: x['dias'], reverse=True)
    top_tiempo_taller = unidades_tiempo[:5]

    choferes_count = Counter()
    for r in reportes:
        ch = obtener_chofer(None, r)
        if ch and ch != 'Personal Operativo':
            choferes_count[ch] += 1
    top_choferes = [{"nombre": k, "cantidad": v} for k, v in choferes_count.most_common(5)]

    unidades_count = Counter()
    for r in reportes:
        u = obtener_unidad(None, r)
        if u:
            unidades_count[u] += 1
    top_unidades_averias = [{"unidad": k, "cantidad": v} for k, v in unidades_count.most_common(5)]

    copes_count = Counter()
    for r in reportes:
        c = r.get('cope') or 'PE CALLE 10'
        copes_count[c.strip()] += 1
    top_copes = [{"cope": k, "cantidad": v} for k, v in copes_count.most_common(5)]

    talleres_dinero = {}
    for f in facturas:
        t = obtener_taller(f, None)
        p = float(f.get('precio_estimado', f.get('precio', 0)) or 0)
        talleres_dinero[t] = talleres_dinero.get(t, 0) + p
    top_talleres_dinero = [{"nombre": k, "total": round(v, 2)} for k, v in sorted(talleres_dinero.items(), key=lambda x: x[1], reverse=True)[:5]]

    # CHARTS
    estados_count = Counter()
    for r in reportes:
        tid = str(r.get('id'))
        fac = next((f for f in facturas if obtener_ticket_id_factura(f) == tid), None)
        if fac:
            st = determinar_estado_ticket(fac)
        else:
            st = "Pendiente de Cotización" if r.get('asignado_a') else "Pendiente / Asig. Taller"
        estados_count[st] += 1

    reportes_ids = {str(r.get('id')) for r in reportes}
    for f in facturas:
        tid = obtener_ticket_id_factura(f)
        if tid not in reportes_ids:
            st = determinar_estado_ticket(f)
            estados_count[st] += 1

    mant_count = Counter()
    for r in reportes:
        m = r.get('mantenimiento') or 'Preventivo'
        mant_count[m.strip().capitalize()] += 1

    copes_chart_count = Counter()
    for r in reportes:
        c = r.get('cope') or 'PE CALLE 10'
        copes_chart_count[c.strip()] += 1

    depto_count = Counter()
    for r in reportes:
        d = r.get('departamento') or 'General'
        depto_count[d.strip()] += 1

    meses_count = Counter()
    for r in reportes:
        f_str = str(r.get('fecha') or r.get('timestamp') or '')
        if len(f_str) >= 7:
            meses_count[f_str[:7]] += 1

    marcas_count = Counter()
    for r in reportes:
        m = r.get('marca') or 'General'
        marcas_count[m.strip().upper()] += 1

    unidades_dinero = {}
    for f in facturas:
        tid = obtener_ticket_id_factura(f)
        r = reportes_map.get(tid)
        u = obtener_unidad(f, r)
        p = float(f.get('precio_estimado', f.get('precio', 0)) or 0)
        unidades_dinero[u] = unidades_dinero.get(u, 0) + p
    top_unidades_dinero = sorted(unidades_dinero.items(), key=lambda x: x[1], reverse=True)[:5]

    from rutas_unidades_metricas import obtener_metricas_flota
    flota_metricas = obtener_metricas_flota()

    return {
        "flota_metricas": flota_metricas,
        "events": {
            "ultimo_llegado": ultimo_llegado,
            "ultimo_finalizado": ultimo_finalizado,
            "ultimo_asignado": ultimo_asignado,
            "ultimo_caro": ultimo_caro,
            "ultimo_reactivado_caras": ultimo_reactivado_caras,
            "ultima_reparada": ultima_reparada,
            "ultima_entregada": ultima_entregada,
            "ultima_rechazada": ultima_rechazada,
            "ultima_fiscal": ultima_fiscal,
            "ultima_garantia": ultima_garantia
        },
        "tops": {
            "top_tiempo_taller": top_tiempo_taller,
            "top_choferes": top_choferes,
            "top_unidades_averias": top_unidades_averias,
            "top_copes": top_copes,
            "top_talleres_dinero": top_talleres_dinero
        },
        "charts": {
            "pie_flota": {
                "labels": ["En Circulación", "En Taller", "Garantía/Retrabajo", "En Jurídicos", "Inhabilitadas (Caras)"],
                "data": [
                    flota_metricas["totales"]["circulacion"],
                    flota_metricas["totales"]["taller"],
                    flota_metricas["totales"].get("garantia", 0),
                    flota_metricas["totales"]["juridicos"],
                    flota_metricas["totales"]["inhabilitadas"]
                ]
            },
            "pie_estados": {"labels": list(estados_count.keys()), "data": list(estados_count.values())},
            "pie_mantenimiento": {"labels": list(mant_count.keys()), "data": list(mant_count.values())},
            "pie_copes": {"labels": list(copes_chart_count.keys()), "data": list(copes_chart_count.values())},
            "pie_deptos": {"labels": list(depto_count.keys()), "data": list(depto_count.values())},
            "bar_meses": {"labels": sorted(list(meses_count.keys())), "data": [meses_count[k] for k in sorted(list(meses_count.keys()))]},
            "bar_talleres": {"labels": [x['nombre'] for x in top_talleres_dinero], "data": [x['total'] for x in top_talleres_dinero]},
            "bar_marcas": {"labels": [k for k, v in marcas_count.most_common(5)], "data": [v for k, v in marcas_count.most_common(5)]},
            "bar_unidades_dinero": {"labels": [k for k, v in top_unidades_dinero], "data": [v for k, v in top_unidades_dinero]}
        }
    }

@analitica_bp.route('/api/analitica/ultimos', methods=['GET'])
def analitica_ultimos():
    period = request.args.get('period', 'Todos')
    datos = obtener_datos_auditados(period)
    return jsonify({
        "status": "success",
        "flota_metricas": datos.get("flota_metricas", {}),
        **datos["events"]
    })

@analitica_bp.route('/api/analitica/flota', methods=['GET'])
def analitica_flota():
    from rutas_unidades_metricas import obtener_metricas_flota
    return jsonify({
        "status": "success",
        **obtener_metricas_flota()
    })

@analitica_bp.route('/api/analitica/tops', methods=['GET'])
def analitica_tops():
    period = request.args.get('period', 'Todos')
    datos = obtener_datos_auditados(period)
    return jsonify({
        "status": "success",
        **datos["tops"]
    })

@analitica_bp.route('/api/analitica/charts', methods=['GET'])
def analitica_charts():
    period = request.args.get('period', 'Todos')
    datos = obtener_datos_auditados(period)
    return jsonify({
        "status": "success",
        **datos["charts"]
    })

@analitica_bp.route('/api/analitica/exportar_excel', methods=['GET'])
def exportar_analitica_excel():
    period = request.args.get('period', 'Todos')
    datos = obtener_datos_auditados(period)
    
    events = datos["events"]
    tops = datos["tops"]
    charts = datos["charts"]
    
    wb = openpyxl.Workbook()
    
    font_title = Font(name='Segoe UI', size=14, bold=True, color='FFFFFF')
    font_subtitle = Font(name='Segoe UI', size=10, italic=True, color='CBD5E1')
    font_header = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    font_section = Font(name='Segoe UI', size=12, bold=True, color='112641')
    font_normal = Font(name='Segoe UI', size=10, color='1E293B')
    font_bold = Font(name='Segoe UI', size=10, bold=True, color='1E293B')

    fill_navy = PatternFill(start_color='112641', end_color='112641', fill_type='solid')
    fill_green = PatternFill(start_color='40916C', end_color='40916C', fill_type='solid')
    fill_sub_header = PatternFill(start_color='1F395A', end_color='1F395A', fill_type='solid')
    fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    # =========================================================
    # HOJA 1: EVENTOS CLAVE
    # =========================================================
    ws1 = wb.active
    ws1.title = "Eventos Clave"
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells('A1:I1')
    ws1['A1'] = "SISTEMA DE FACTURACION CENTRAL - ANALÍTICA GLOBAL"
    ws1['A1'].font = font_title
    ws1['A1'].fill = fill_navy
    ws1['A1'].alignment = align_center
    ws1.row_dimensions[1].height = 36

    ws1.merge_cells('A2:I2')
    ws1['A2'] = f"Reporte de Eventos Clave | Periodo: {period} | Generado: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws1['A2'].font = font_subtitle
    ws1['A2'].fill = fill_navy
    ws1['A2'].alignment = align_center
    ws1.row_dimensions[2].height = 22

    headers_ev = ["Evento Clave", "Folio / Ticket", "Chofer / Conductor", "Unidad (Eco)", "Taller Asignado", "Administrador", "Supervisor", "Jefatura", "Fecha / Registro"]
    for col_idx, h in enumerate(headers_ev, 1):
        cell = ws1.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_green
        cell.alignment = align_center
        cell.border = border_thin
    ws1.row_dimensions[4].height = 28

    event_labels = [
        ("Último Reporte Llegado", events.get("ultimo_llegado")),
        ("Último Reporte Finalizado", events.get("ultimo_finalizado")),
        ("Último Asignado a Taller", events.get("ultimo_asignado")),
        ("Última Cotización Cara (>=10k)", events.get("ultimo_caro")),
        ("Último Reactivado Cot. Caras", events.get("ultimo_reactivado_caras")),
        ("Última Unidad Reparada", events.get("ultima_reparada")),
        ("Última Unidad Entregada", events.get("ultima_entregada")),
        ("Última Cotización Rechazada", events.get("ultima_rechazada")),
        ("Última Factura Fiscal", events.get("ultima_fiscal")),
        ("Última Unidad a Garantía / Retrabajo", events.get("ultima_garantia"))
    ]

    row_num = 5
    for label, ev in event_labels:
        row_vals = [
            label,
            ev.get("ticket", "") if ev else "",
            ev.get("chofer", "") if ev else "",
            ev.get("unidad", "") if ev else "",
            ev.get("taller", "") if ev else "",
            ev.get("administrador", "") if ev else "",
            ev.get("supervisor", "") if ev else "",
            ev.get("jefatura", "") if ev else "",
            str(ev.get("fecha", ""))[:19].replace('T', ' ') if ev else ""
        ]
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws1.cell(row=row_num, column=col_idx, value=val)
            cell.font = font_normal
            cell.border = border_thin
            if col_idx in [2, 4, 9]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
            if row_num % 2 == 0:
                cell.fill = fill_zebra
        ws1.row_dimensions[row_num].height = 22
        row_num += 1

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 15)

    # =========================================================
    # HOJA 2: RANKINGS OPERATIVOS
    # =========================================================
    ws2 = wb.create_sheet(title="Rankings y Tops")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells('A1:G1')
    ws2['A1'] = "INDICADORES CLAVE DE DESEMPEÑO (RANKINGS / TOPS)"
    ws2['A1'].font = font_title
    ws2['A1'].fill = fill_navy
    ws2['A1'].alignment = align_center
    ws2.row_dimensions[1].height = 36

    ws2.merge_cells('A2:G2')
    ws2['A2'] = f"Periodo: {period} | Sistema de Facturacion Central"
    ws2['A2'].font = font_subtitle
    ws2['A2'].fill = fill_navy
    ws2['A2'].alignment = align_center
    ws2.row_dimensions[2].height = 22

    current_r = 4

    def escribir_seccion_top(ws, start_row, titulo, headers, data_rows):
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=len(headers))
        title_cell = ws.cell(row=start_row, column=1, value=titulo)
        title_cell.font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        title_cell.fill = fill_sub_header
        title_cell.alignment = align_left
        ws.row_dimensions[start_row].height = 24
        
        h_row = start_row + 1
        for c_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=h_row, column=c_idx, value=h)
            cell.font = font_header
            cell.fill = fill_green
            cell.alignment = align_center
            cell.border = border_thin
        ws.row_dimensions[h_row].height = 24
        
        r_idx = h_row + 1
        if not data_rows:
            ws.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=len(headers))
            empty_cell = ws.cell(row=r_idx, column=1, value="Sin datos registrados en el periodo")
            empty_cell.font = Font(name='Segoe UI', size=9, italic=True, color='64748B')
            empty_cell.alignment = align_center
            empty_cell.border = border_thin
            ws.row_dimensions[r_idx].height = 20
            return r_idx + 2
            
        for row in data_rows:
            for c_idx, val in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = font_normal
                cell.border = border_thin
                if isinstance(val, (int, float)):
                    cell.alignment = align_right
                    if c_idx == len(row) and "Facturación" in titulo:
                        cell.number_format = '$#,##0.00'
                elif c_idx == 1:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left
                if r_idx % 2 == 0:
                    cell.fill = fill_zebra
            ws.row_dimensions[r_idx].height = 20
            r_idx += 1
            
        return r_idx + 1

    # Top 1: Tiempo en taller
    data_taller = [[f"#{i+1}", f"Eco #{item.get('unidad')}", f"{item.get('dias')} días", item.get('taller')] for i, item in enumerate(tops["top_tiempo_taller"])]
    current_r = escribir_seccion_top(ws2, current_r, "⏳ Top 5 Unidades con Mayor Tiempo en Taller", ["Posición", "Unidad Vehicular", "Días en Taller", "Taller Asignado"], data_taller)

    # Top 2: Unidades Averias
    data_averias = [[f"#{i+1}", f"Eco #{item.get('unidad')}", item.get('cantidad')] for i, item in enumerate(tops["top_unidades_averias"])]
    current_r = escribir_seccion_top(ws2, current_r, "⚠️ Top 5 Unidades con Más Averías", ["Posición", "Unidad Vehicular", "Total de Reportes"], data_averias)

    # Top 3: Choferes
    data_choferes = [[f"#{i+1}", item.get('nombre'), item.get('cantidad')] for i, item in enumerate(tops["top_choferes"])]
    current_r = escribir_seccion_top(ws2, current_r, "👨‍✈️ Top 5 Choferes con Más Reportes", ["Posición", "Nombre del Chofer", "Total de Reportes"], data_choferes)

    # Top 4: Talleres Dinero
    data_dinero = [[f"#{i+1}", item.get('nombre'), item.get('total')] for i, item in enumerate(tops["top_talleres_dinero"])]
    current_r = escribir_seccion_top(ws2, current_r, "💰 Top 5 Talleres con Mayor Facturación", ["Posición", "Taller / Proveedor", "Facturación Total ($)"], data_dinero)

    # Top 5: COPEs
    data_copes = [[f"#{i+1}", item.get('cope'), item.get('cantidad')] for i, item in enumerate(tops["top_copes"])]
    current_r = escribir_seccion_top(ws2, current_r, "🏢 Top 5 COPEs con Mayor Actividad", ["Posición", "COPE Operativo", "Total de Reportes"], data_copes)

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 4, 18)

    # =========================================================
    # HOJA 3: DISTRIBUCIÓN Y MÉTRICAS
    # =========================================================
    ws3 = wb.create_sheet(title="Distribución y Gráficas")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells('A1:F1')
    ws3['A1'] = "MÉTRICAS DE DISTRIBUCIÓN Y VOLUMEN OPERATIVO"
    ws3['A1'].font = font_title
    ws3['A1'].fill = fill_navy
    ws3['A1'].alignment = align_center
    ws3.row_dimensions[1].height = 36

    ws3.merge_cells('A2:F2')
    ws3['A2'] = f"Datos consolidados de gráficas analíticas | Periodo: {period}"
    ws3['A2'].font = font_subtitle
    ws3['A2'].fill = fill_navy
    ws3['A2'].alignment = align_center
    ws3.row_dimensions[2].height = 22

    cur_r3 = 4

    def escribir_tabla_metricas(ws, start_row, titulo, h1, h2, labels, data, es_moneda=False):
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=2)
        title_cell = ws.cell(row=start_row, column=1, value=titulo)
        title_cell.font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        title_cell.fill = fill_sub_header
        title_cell.alignment = align_left
        ws.row_dimensions[start_row].height = 24
        
        h_row = start_row + 1
        c1 = ws.cell(row=h_row, column=1, value=h1)
        c2 = ws.cell(row=h_row, column=2, value=h2)
        for c in [c1, c2]:
            c.font = font_header
            c.fill = fill_green
            c.alignment = align_center
            c.border = border_thin
        ws.row_dimensions[h_row].height = 24
        
        r_idx = h_row + 1
        for lbl, val in zip(labels, data):
            cell1 = ws.cell(row=r_idx, column=1, value=str(lbl))
            cell2 = ws.cell(row=r_idx, column=2, value=val)
            for cell in [cell1, cell2]:
                cell.font = font_normal
                cell.border = border_thin
                if r_idx % 2 == 0:
                    cell.fill = fill_zebra
            cell1.alignment = align_left
            cell2.alignment = align_right
            if es_moneda:
                cell2.number_format = '$#,##0.00'
            ws.row_dimensions[r_idx].height = 20
            r_idx += 1
            
        return r_idx + 1

    if "pie_flota" in charts:
        cur_r3 = escribir_tabla_metricas(ws3, cur_r3, "🚗 Estado Operativo de la Flota Vehicular", "Estado Operativo", "Cantidad de Unidades", charts["pie_flota"]["labels"], charts["pie_flota"]["data"])
    cur_r3 = escribir_tabla_metricas(ws3, cur_r3, "📊 Distribución por Estado de Tickets", "Estado del Ticket", "Cantidad de Tickets", charts["pie_estados"]["labels"], charts["pie_estados"]["data"])
    cur_r3 = escribir_tabla_metricas(ws3, cur_r3, "🔧 Tipos de Mantenimiento", "Tipo de Mantenimiento", "Cantidad de Tickets", charts["pie_mantenimiento"]["labels"], charts["pie_mantenimiento"]["data"])
    cur_r3 = escribir_tabla_metricas(ws3, cur_r3, "🏢 Distribución por COPE", "COPE", "Cantidad de Tickets", charts["pie_copes"]["labels"], charts["pie_copes"]["data"])
    cur_r3 = escribir_tabla_metricas(ws3, cur_r3, "🏢 Fallas por Departamento", "Departamento", "Cantidad de Fallas", charts["pie_deptos"]["labels"], charts["pie_deptos"]["data"])
    cur_r3 = escribir_tabla_metricas(ws3, cur_r3, "📅 Tickets Generados por Mes", "Mes (YYYY-MM)", "Cantidad de Tickets", charts["bar_meses"]["labels"], charts["bar_meses"]["data"])
    cur_r3 = escribir_tabla_metricas(ws3, cur_r3, "🚗 Modelos de Unidad (+ Fallas)", "Marca / Modelo", "Cantidad de Fallas", charts["bar_marcas"]["labels"], charts["bar_marcas"]["data"])
    cur_r3 = escribir_tabla_metricas(ws3, cur_r3, "💸 Top 5 Unidades (+ Gasto Acumulado)", "Unidad (Eco)", "Gasto Total ($)", [f"Eco #{u}" for u in charts["bar_unidades_dinero"]["labels"]], charts["bar_unidades_dinero"]["data"], es_moneda=True)

    for col in ws3.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = max(max_len + 5, 22)

    # Guardar en buffer en memoria
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    fecha_archivo = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    period_clean = period.replace('-', '_').replace(' ', '_')
    filename = f"Analitica_Facturacion_Central_{period_clean}_{fecha_archivo}.xlsx"
    
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )
